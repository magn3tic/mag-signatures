"""Deletes the pages of everyone marked Remove in the master sheet.

    python3 generator/remove.py            # show what would go
    python3 generator/remove.py --apply    # delete them and rebuild the index

The headshots are deliberately left in images/. Emails already sent still
hotlink them, and pulling the file breaks the signature in every one of them.
"""
import os, subprocess, sys
from openpyxl import load_workbook

XLSX = "MAG-signatures-master.xlsx"
apply_ = "--apply" in sys.argv

if not os.path.exists(XLSX):
    sys.exit("No %s -- run generator/build-master.py first." % XLSX)

ws = load_workbook(XLSX, data_only=True)["Signatures"]
head = [c.value for c in ws[1]]
i_status, i_name, i_page = head.index("Status"), head.index("Name"), head.index("Page")

doomed = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if str(row[i_status] or "").strip().lower() == "remove":
        doomed.append((row[i_name], row[i_page]))

if not doomed:
    print("Nobody is marked Remove.")
    sys.exit(0)

print("%d marked Remove:" % len(doomed))
missing = []
for name, page in doomed:
    path = os.path.join("pages", page)
    here = os.path.exists(path)
    if not here:
        missing.append(page)
    print("   %-28s %s%s" % (name, page, "" if here else "   (already gone)"))

if not apply_:
    print("\nNothing deleted. Re-run with --apply to delete these and rebuild the index.")
    sys.exit(0)

for _, page in doomed:
    path = os.path.join("pages", page)
    if os.path.exists(path):
        os.remove(path)

# The index and people.json are both derived from whatever is left in pages/.
for script in ("extract.py", "build-index.py", "build-master.py"):
    subprocess.run([sys.executable, os.path.join("generator", script)], check=True)

print("\nDone. Headshots left in images/ on purpose -- already-sent emails still load them.")
